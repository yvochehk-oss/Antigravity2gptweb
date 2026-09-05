"""Canonical/Statutory multi-sheet Excel export service.

The exporter is a read-only reporting boundary. Business facts come from
``analytics_canonical_facts_current`` / existing Canonical read models, while
statutory VAT comes only from the current SUCCEEDED formal VAT ledger. Missing
formal resources remain blank and are never coerced to zero.
"""
from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
from io import BytesIO
import json
import re
from typing import Any, Iterable
from urllib.parse import quote

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import inspect, select, text
from sqlalchemy.exc import SQLAlchemyError

from app.models import Project
from app.services.canonical_project_summary import canonical_project_summary
from app.services.canonical_ssot import (
    DEDUCTIBILITY_ELIGIBLE,
    DEDUCTIBILITY_INELIGIBLE,
    invoice_deductibility_status,
)
from app.services.canonical_v3_bridge import CanonicalV3Bridge
from app.services.formal_vat_statutory import (
    FormalVatStatutoryResourceIntegrityError,
    FormalVatStatutoryResourceNotFoundError,
    get_formal_vat_statutory_resource,
)
from app.services.legal_entity_master_data import list_legal_entities
from app.services.legal_entity_scope import aggregate_legal_entity_scope

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MONEY_FORMAT = "#,##0.00"
PERCENT_FORMAT = "0.00%"
CANONICAL_VIEW = "analytics_canonical_facts_current"
_PERIOD_RE = re.compile(r"^\d{4}-(?:0[1-9]|1[0-2])$")

SHEET_GROUP = "集团总览"
SHEET_ENTITIES = "法人主体分户台账"
SHEET_PROJECTS = "项目工程库分明细台账"
SHEET_OUTPUT = "收入发票交易明细"
SHEET_INPUT = "支出发票与成本明细"
SHEET_PREPAY = "税款预缴与完税凭证明细"
SHEET_NAMES = [SHEET_GROUP, SHEET_ENTITIES, SHEET_PROJECTS, SHEET_OUTPUT, SHEET_INPUT, SHEET_PREPAY]

_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_SECTION_FILL = PatternFill("solid", fgColor="D9EAF7")
_SUBTLE_FILL = PatternFill("solid", fgColor="F3F6F9")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_TITLE_FONT = Font(size=16, bold=True)
_BOLD_FONT = Font(bold=True)
_THIN = Side(style="thin", color="D9E1F2")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


@dataclass(frozen=True)
class ExcelExportArtifact:
    content: bytes
    filename: str
    media_type: str = XLSX_MEDIA_TYPE


def _money(value: Any, *, none_if_missing: bool = False) -> Decimal | None:
    if value is None and none_if_missing:
        return None
    try:
        number = Decimal(str(value if value is not None else 0))
    except (InvalidOperation, TypeError, ValueError):
        return None if none_if_missing else Decimal("0.00")
    if not number.is_finite():
        return None if none_if_missing else Decimal("0.00")
    return number.quantize(Decimal("0.01"))


def _payload(value: Any) -> dict[str, Any]:
    if isinstance(value, dict):
        return dict(value)
    if isinstance(value, str):
        try:
            parsed = json.loads(value)
        except json.JSONDecodeError:
            return {}
        return parsed if isinstance(parsed, dict) else {}
    return {}


def _text(value: Any) -> str:
    return " ".join(str(value or "").strip().split())


def _code(value: Any) -> str:
    return _text(value).upper()


def _validate_period(period: str | None, *, required: bool) -> str | None:
    value = str(period or "").strip()
    if not value:
        if required:
            raise ValueError("period is required for current view")
        return None
    if not _PERIOD_RE.fullmatch(value):
        raise ValueError("period must be YYYY-MM")
    return value


def _fact_period(payload: dict[str, Any]) -> str:
    raw = _text(
        payload.get("invoice_date")
        or payload.get("transaction_date")
        or payload.get("date")
        or payload.get("period")
    )
    return raw[:7] if len(raw) >= 7 and raw[4] == "-" else ""


def _fact_date(payload: dict[str, Any]) -> str:
    raw = _text(payload.get("invoice_date") or payload.get("transaction_date") or payload.get("date"))
    return raw[:10] if len(raw) >= 10 else raw


def _canonical_available(db) -> bool:
    try:
        return bool(inspect(db.get_bind()).has_table(CANONICAL_VIEW))
    except (SQLAlchemyError, AttributeError):
        return False


def _legal_master(db, scope: str) -> tuple[list[dict[str, Any]], dict[str, str]]:
    all_entities = list_legal_entities(db, active=True, legal_entity=True)
    by_code = {str(row["canonical_code"]).upper(): row for row in all_entities}
    wanted = _code(scope or "ALL") or "ALL"
    if wanted == "ALL":
        selected = all_entities
    else:
        row = by_code.get(wanted)
        if row is None:
            raise ValueError(f"unknown legal entity scope: {wanted}")
        selected = [row]
    names = {str(row["canonical_code"]).upper(): str(row["legal_name"]) for row in all_entities}
    return selected, names


def _official_periods(db, entity_code: str, *, view: str, period: str | None) -> list[str]:
    where = [
        "tps.tax_type = 'VAT'",
        "tps.current_run_id IS NOT NULL",
        "UPPER(ie.canonical_code) = UPPER(:entity_code)",
    ]
    params: dict[str, Any] = {"entity_code": entity_code}
    if view == "current":
        where.append("tps.tax_period = CAST(:period || '-01' AS DATE)")
        params["period"] = period
    elif period:
        where.append("tps.tax_period <= CAST(:period || '-01' AS DATE)")
        params["period"] = period
    rows = db.execute(
        text(
            "SELECT to_char(tps.tax_period, 'YYYY-MM') AS period "
            "FROM tax_period_states tps "
            "JOIN internal_entities ie ON ie.party_id = tps.reporting_party_id "
            f"WHERE {' AND '.join(where)} ORDER BY tps.tax_period"
        ),
        params,
    ).scalars().all()
    return [str(item) for item in rows if item]


def _formal_aggregate(db, entity_code: str, *, view: str, period: str | None) -> dict[str, Any]:
    periods = _official_periods(db, entity_code, view=view, period=period)
    resources: list[dict[str, Any]] = []
    integrity_error = ""
    for item_period in periods:
        try:
            resources.append(get_formal_vat_statutory_resource(db, entity_code, item_period))
        except FormalVatStatutoryResourceNotFoundError:
            continue
        except FormalVatStatutoryResourceIntegrityError as exc:
            integrity_error = str(exc)
            break

    if integrity_error:
        return {
            "status": "FORMAL_INVALID",
            "periods": periods,
            "opening_input_credit": None,
            "output_vat": None,
            "input_vat": None,
            "tax_prepayment": None,
            "vat_payable_after_prepayment": None,
            "closing_input_credit": None,
            "detail": integrity_error,
        }
    if not resources:
        return {
            "status": "FORMAL_NOT_AVAILABLE",
            "periods": periods,
            "opening_input_credit": None,
            "output_vat": None,
            "input_vat": None,
            "tax_prepayment": None,
            "vat_payable_after_prepayment": None,
            "closing_input_credit": None,
            "detail": "",
        }

    ledgers = [item["vat_ledger"] for item in resources]
    return {
        "status": "FORMAL_READY" if len(resources) == len(periods) else "FORMAL_PARTIAL",
        "periods": [item["period"] for item in resources],
        "opening_input_credit": _money(ledgers[0]["opening_input_credit"]),
        "output_vat": sum((_money(row["output_vat"]) or Decimal("0")) for row in ledgers),
        "input_vat": sum((_money(row["input_vat"]) or Decimal("0")) for row in ledgers),
        "tax_prepayment": sum((_money(row["tax_prepayment"]) or Decimal("0")) for row in ledgers),
        "vat_payable_after_prepayment": sum(
            (_money(row["vat_payable_after_prepayment"]) or Decimal("0")) for row in ledgers
        ),
        "closing_input_credit": _money(ledgers[-1]["closing_input_credit"]),
        "detail": "",
    }


def _projection(db, entity_code: str, *, view: str, period: str | None) -> dict[str, Any]:
    try:
        data = aggregate_legal_entity_scope(db, entity_code, period=period)
    except (SQLAlchemyError, ValueError):
        return {
            "status": "PROJECTION_UNAVAILABLE",
            "revenue": None,
            "book_cost_projection": None,
            "accounting_profit_projection": None,
            "deductible_input_vat": None,
            "nondeductible_input_vat": None,
            "pending_input_vat": None,
            "project_contributions": [],
        }

    cumulative = data.get("cumulative") if isinstance(data.get("cumulative"), dict) else {}
    basis = cumulative if view == "cumulative" else data
    return {
        "status": str(data.get("status") or "READY"),
        "revenue": _money(basis.get("revenue"), none_if_missing=True),
        "book_cost_projection": _money(basis.get("book_cost_projection"), none_if_missing=True),
        "accounting_profit_projection": _money(
            basis.get("accounting_profit_projection"), none_if_missing=True
        ),
        # Deductibility buckets are exact-period in current view. With cumulative
        # view they are only populated when no cutoff is supplied (whole dataset).
        "deductible_input_vat": _money(
            data.get("deductible_input_vat"), none_if_missing=True
        ) if view == "current" or period is None else None,
        "nondeductible_input_vat": _money(
            data.get("nondeductible_input_vat"), none_if_missing=True
        ) if view == "current" or period is None else None,
        "pending_input_vat": _money(
            data.get("pending_input_vat"), none_if_missing=True
        ) if view == "current" or period is None else None,
        "project_contributions": list(data.get("project_contributions") or []),
    }


def _invoice_rows(
    db,
    *,
    selected_codes: set[str],
    names: dict[str, str],
    view: str,
    period: str | None,
    projects: dict[int, Project],
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    if not _canonical_available(db):
        return [], []
    try:
        raw_rows = db.execute(
            text(
                "SELECT fact_id, project_id, business_key, fact_version, payload "
                "FROM analytics_canonical_facts_current "
                "WHERE lower(fact_type) = 'invoice' "
                "ORDER BY project_id NULLS LAST, business_key, fact_version, fact_id"
            )
        ).mappings().all()
    except SQLAlchemyError:
        return [], []

    output: list[dict[str, Any]] = []
    input_rows: list[dict[str, Any]] = []
    for raw in raw_rows:
        payload = _payload(raw["payload"])
        item_period = _fact_period(payload)
        if view == "current" and item_period != period:
            continue
        if view == "cumulative" and period and (not item_period or item_period > period):
            continue
        seller = _code(payload.get("seller_entity_code") or payload.get("seller_code"))
        buyer = _code(payload.get("buyer_entity_code") or payload.get("buyer_code"))
        if seller not in selected_codes and buyer not in selected_codes:
            continue

        project_id = int(raw["project_id"]) if raw["project_id"] is not None else None
        project = projects.get(project_id) if project_id is not None else None
        vat = _money(payload.get("vat_amount") if payload.get("vat_amount") is not None else payload.get("tax_amount")) or Decimal("0")
        net_value = payload.get("net_amount")
        gross_value = payload.get("gross_amount")
        if gross_value is None:
            gross_value = payload.get("amount_with_tax")
        if gross_value is None:
            gross_value = payload.get("total_amount")
        if gross_value is None:
            gross_value = payload.get("amount")
        gross = _money(gross_value, none_if_missing=True)
        net = _money(net_value, none_if_missing=True)
        if net is None and gross is not None:
            net = (gross - vat).quantize(Decimal("0.01"))
        if gross is None and net is not None:
            gross = (net + vat).quantize(Decimal("0.01"))
        contract_no = _text(
            payload.get("contract_no") or payload.get("contract_code") or payload.get("contract_number")
        )
        invoice_no = _text(
            payload.get("invoice_no") or payload.get("invoice_number") or raw["business_key"]
        )
        summary = _text(
            payload.get("service_summary")
            or payload.get("goods_or_service")
            or payload.get("item_name")
            or payload.get("description")
            or payload.get("summary")
        )
        common = {
            "invoice_date": _fact_date(payload),
            "period": item_period,
            "project_id": project_id,
            "project_code": _text(project.project_code or project.code) if project else _text(payload.get("project_code")),
            "project_name": _text(project.name) if project else _text(payload.get("project_name")),
            "contract_no": contract_no,
            "invoice_no": invoice_no,
            "summary": summary,
            "net_amount": net,
            "vat_amount": vat,
            "gross_amount": gross,
            "fact_status": "CURRENT_CANONICAL",
            "fact_id": int(raw["fact_id"]) if raw["fact_id"] is not None else None,
        }
        if seller in selected_codes:
            output.append({
                **common,
                "entity_code": seller,
                "entity_name": names.get(seller, seller),
                "counterparty_name": _text(
                    payload.get("buyer_name") or payload.get("buyer_legal_name") or buyer
                ),
            })
        if buyer in selected_codes:
            deductibility = invoice_deductibility_status(payload)
            if deductibility == DEDUCTIBILITY_ELIGIBLE:
                deductible_label = "可抵扣"
            elif deductibility == DEDUCTIBILITY_INELIGIBLE:
                deductible_label = "不可抵扣"
            else:
                deductible_label = "待确认"
            input_rows.append({
                **common,
                "entity_code": buyer,
                "entity_name": names.get(buyer, buyer),
                "counterparty_name": _text(
                    payload.get("seller_name") or payload.get("seller_legal_name") or seller
                ),
                "deductibility_status": deductible_label,
            })
    return output, input_rows


def _prepayment_rows(
    db,
    *,
    selected_codes: set[str],
    names: dict[str, str],
    view: str,
    period: str | None,
) -> list[dict[str, Any]]:
    where = [
        "tpf.tax_type = 'VAT'",
        "f.is_current = TRUE",
        "f.validation_status = 'VALID'",
    ]
    params: dict[str, Any] = {}
    if selected_codes:
        where.append("UPPER(ie.canonical_code) = ANY(:entity_codes)")
        params["entity_codes"] = sorted(selected_codes)
    if view == "current":
        where.append("tpf.tax_period = CAST(:period || '-01' AS DATE)")
        params["period"] = period
    elif period:
        where.append("tpf.tax_period <= CAST(:period || '-01' AS DATE)")
        params["period"] = period
    try:
        rows = db.execute(
            text(
                "SELECT tpf.tax_event_date, tpf.tax_period, tpf.tax_type, tpf.tax_amount, "
                "tpf.external_reference, tpf.note, tpf.project_id, ie.canonical_code, "
                "p.name AS entity_name, pr.code AS project_code "
                "FROM tax_prepayment_facts tpf "
                "JOIN facts f ON f.id = tpf.fact_id "
                "JOIN internal_entities ie ON ie.party_id = tpf.reporting_party_id "
                "JOIN parties p ON p.id = ie.party_id "
                "LEFT JOIN projects pr ON pr.id = tpf.project_id "
                f"WHERE {' AND '.join(where)} "
                "ORDER BY tpf.tax_event_date, tpf.fact_id"
            ),
            params,
        ).mappings().all()
    except SQLAlchemyError:
        return []
    return [
        {
            "prepayment_date": str(row["tax_event_date"] or ""),
            "entity_code": _code(row["canonical_code"]),
            "entity_name": _text(row["entity_name"] or names.get(_code(row["canonical_code"]), "")),
            "project_code": _text(row["project_code"]),
            "receipt_no": _text(row["external_reference"]),
            "tax_authority": "",
            "tax_type": _text(row["tax_type"]),
            "prepayment_amount": _money(row["tax_amount"]),
            "note": _text(row["note"]),
        }
        for row in rows
    ]


def _canonical_four_flow_percentage(db, project_id: int) -> tuple[Decimal | None, str]:
    try:
        flow = CanonicalV3Bridge(db).four_flow(project_id)
    except Exception:
        return None, "UNAVAILABLE"
    relations = list(flow.get("relationships") or [])
    if not relations:
        return None, "UNAVAILABLE"
    # Canonical bridge currently models contract/invoice/payment deterministic
    # business-key links. Fulfilment evidence has no Canonical contract-key
    # relation yet, so the percentage intentionally measures the three proven
    # canonical links instead of fabricating a fourth flow.
    expected = len(relations) * 3
    available = 0
    for row in relations:
        available += int(bool(row.get("contract_fact_ids")))
        available += int(bool(row.get("invoice_fact_ids")))
        available += int(bool(row.get("payment_fact_ids")))
    percentage = (Decimal(available) / Decimal(expected) * Decimal("100")).quantize(Decimal("0.01"))
    return percentage, "CANONICAL_3_FLOW_LINKAGE"


def _project_rows(
    db,
    *,
    selected_codes: set[str],
    names: dict[str, str],
    invoice_output: list[dict[str, Any]],
    invoice_input: list[dict[str, Any]],
    projection_by_entity: dict[str, dict[str, Any]],
) -> list[dict[str, Any]]:
    projects = db.scalars(select(Project).order_by(Project.id)).all()
    contribution_codes: dict[int, set[str]] = {}
    for entity_code, projection in projection_by_entity.items():
        for item in projection.get("project_contributions") or []:
            project_id = item.get("project_id")
            if project_id is not None:
                contribution_codes.setdefault(int(project_id), set()).add(entity_code)

    output_by_project: dict[int, Decimal] = {}
    input_by_project: dict[int, Decimal] = {}
    for row in invoice_output:
        if row["project_id"] is not None:
            output_by_project[int(row["project_id"])] = output_by_project.get(int(row["project_id"]), Decimal("0")) + (row["gross_amount"] or Decimal("0"))
    for row in invoice_input:
        if row["project_id"] is not None:
            input_by_project[int(row["project_id"])] = input_by_project.get(int(row["project_id"]), Decimal("0")) + (row["gross_amount"] or Decimal("0"))

    result: list[dict[str, Any]] = []
    for project in projects:
        entity_code = _code(project.entity_code)
        derived_codes = contribution_codes.get(int(project.id), set())
        if not entity_code and len(derived_codes) == 1:
            entity_code = next(iter(derived_codes))
        if selected_codes and entity_code and entity_code not in selected_codes:
            continue
        if selected_codes and not entity_code and not (derived_codes & selected_codes):
            continue
        try:
            summary = canonical_project_summary(db, int(project.id))
            contract_total = _money(summary.get("contract_total"), none_if_missing=True)
            real_cost = _money(summary.get("real_cost"), none_if_missing=True)
            paid_amount = _money(summary.get("external_cash_cost"), none_if_missing=True)
            summary_status = "READY"
        except Exception:
            contract_total = None
            real_cost = None
            paid_amount = None
            summary_status = "DEGRADED"
        completeness, completeness_source = _canonical_four_flow_percentage(db, int(project.id))
        if completeness is None:
            risk = "未知"
        elif completeness < Decimal("75"):
            risk = "高"
        elif completeness < Decimal("100"):
            risk = "中"
        else:
            risk = "低"
        result.append({
            "project_code": _text(project.project_code or project.code),
            "project_name": _text(project.name),
            "entity_code": entity_code,
            "entity_name": names.get(entity_code, entity_code),
            "contract_total": contract_total,
            "real_cost": real_cost,
            "output_invoice_amount": output_by_project.get(int(project.id), Decimal("0.00")),
            "input_invoice_amount": input_by_project.get(int(project.id), Decimal("0.00")),
            "paid_amount": paid_amount,
            "four_flow_completeness": completeness,
            "four_flow_source": completeness_source,
            "risk_rating": risk,
            "status": summary_status,
        })
    return result


def collect_excel_export_data(
    db,
    *,
    scope: str = "ALL",
    view: str = "current",
    period: str | None = None,
) -> dict[str, Any]:
    """Collect one immutable read snapshot for the workbook."""
    normalized_view = str(view or "").strip().lower()
    if normalized_view not in {"current", "cumulative"}:
        raise ValueError("view must be current or cumulative")
    normalized_period = _validate_period(period, required=normalized_view == "current")
    selected, names = _legal_master(db, scope)
    selected_codes = {str(row["canonical_code"]).upper() for row in selected}
    projects = {int(item.id): item for item in db.scalars(select(Project).order_by(Project.id)).all()}

    entity_rows: list[dict[str, Any]] = []
    projection_by_entity: dict[str, dict[str, Any]] = {}
    for entity in selected:
        code = str(entity["canonical_code"]).upper()
        projection = _projection(db, code, view=normalized_view, period=normalized_period)
        projection_by_entity[code] = projection
        formal = _formal_aggregate(db, code, view=normalized_view, period=normalized_period)
        period_label = normalized_period if normalized_view == "current" else (
            f"累计至{normalized_period}" if normalized_period else "累计全部期间"
        )
        entity_rows.append({
            "entity_code": code,
            "entity_name": str(entity["legal_name"]),
            "business_role": code[:1],
            "period": period_label,
            "opening_input_credit": formal["opening_input_credit"],
            "output_vat": formal["output_vat"],
            "input_vat": formal["input_vat"],
            "deductible_input_vat": projection["deductible_input_vat"],
            "nondeductible_input_vat": projection["nondeductible_input_vat"],
            "pending_input_vat": projection["pending_input_vat"],
            "tax_prepayment": formal["tax_prepayment"],
            "vat_payable": formal["vat_payable_after_prepayment"],
            "closing_input_credit": formal["closing_input_credit"],
            "revenue_projection": projection["revenue"],
            "book_cost_projection": projection["book_cost_projection"],
            "accounting_profit_projection": projection["accounting_profit_projection"],
            "gate_status": formal["status"],
            "projection_status": projection["status"],
        })

    output_rows, input_rows = _invoice_rows(
        db,
        selected_codes=selected_codes,
        names=names,
        view=normalized_view,
        period=normalized_period,
        projects=projects,
    )
    prepayment_rows = _prepayment_rows(
        db,
        selected_codes=selected_codes,
        names=names,
        view=normalized_view,
        period=normalized_period,
    )
    project_rows = _project_rows(
        db,
        selected_codes=selected_codes,
        names=names,
        invoice_output=output_rows,
        invoice_input=input_rows,
        projection_by_entity=projection_by_entity,
    )
    return {
        "metadata": {
            "scope": _code(scope or "ALL") or "ALL",
            "view": normalized_view,
            "period": normalized_period or "全部期间",
            "generated_at": datetime.now(timezone.utc).isoformat(),
            "business_fact_source": CANONICAL_VIEW,
            "statutory_source": "tax_period_states.current_run_id->calculation_runs->entity_vat_ledgers",
            "projection_source": "TAX_ENGINE / LEGAL_ENTITY_PROJECTION (non-filing-basis)",
            "legal_entity_count": len(entity_rows),
        },
        "entities": entity_rows,
        "projects": project_rows,
        "output_invoices": output_rows,
        "input_invoices": input_rows,
        "prepayments": prepayment_rows,
    }


def _write_table(
    ws,
    *,
    headers: list[str],
    rows: Iterable[Iterable[Any]],
    money_columns: set[int] | None = None,
    percent_columns: set[int] | None = None,
    start_row: int = 1,
) -> int:
    money_columns = money_columns or set()
    percent_columns = percent_columns or set()
    header_row = start_row
    for col, value in enumerate(headers, start=1):
        cell = ws.cell(header_row, col, value)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _BORDER
    row_no = header_row + 1
    for values in rows:
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row_no, col, value)
            cell.border = _BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=col not in money_columns)
            if col in money_columns and value is not None:
                cell.number_format = MONEY_FORMAT
            if col in percent_columns and value is not None:
                # Stored completeness values are 0..100, not fractional.
                cell.number_format = "0.00"
        row_no += 1
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(headers))}{max(header_row, row_no - 1)}"
    ws.freeze_panes = f"A{header_row + 1}"
    return row_no


def _autosize(ws, *, min_width: int = 10, max_width: int = 42) -> None:
    for column_cells in ws.columns:
        letter = get_column_letter(column_cells[0].column)
        width = min_width
        for cell in column_cells[:300]:
            value = "" if cell.value is None else str(cell.value)
            display_len = sum(2 if ord(char) > 127 else 1 for char in value)
            width = max(width, min(display_len + 2, max_width))
        ws.column_dimensions[letter].width = width


def build_excel_workbook(data: dict[str, Any]) -> bytes:
    """Render six professional worksheets from an already-collected snapshot."""
    wb = Workbook()
    default = wb.active
    wb.remove(default)
    meta = data["metadata"]
    entities = list(data.get("entities") or [])

    ws = wb.create_sheet(SHEET_GROUP)
    ws["A1"] = "集团税务与经营 Excel 总览"
    ws["A1"].font = _TITLE_FONT
    ws.merge_cells("A1:H1")
    metadata_rows = [
        ("数据范围", meta["scope"]),
        ("视图", meta["view"]),
        ("所属期间", meta["period"]),
        ("法人主体数量", meta["legal_entity_count"]),
        ("权威业务事实来源", meta["business_fact_source"]),
        ("法定 VAT 来源", meta["statutory_source"]),
        ("经营投影来源", meta["projection_source"]),
        ("生成时间 UTC", meta["generated_at"]),
    ]
    for idx, (label, value) in enumerate(metadata_rows, start=3):
        ws.cell(idx, 1, label).font = _BOLD_FONT
        ws.cell(idx, 1).fill = _SECTION_FILL
        ws.cell(idx, 2, value)
    overview_headers = [
        "法人编码", "法人名称", "期间", "销项税额", "进项税额", "实际应纳增值税",
        "期末留抵", "经营投影收入", "账面成本投影", "会计利润投影", "门禁状态",
    ]
    overview_rows = [
        [
            row["entity_code"], row["entity_name"], row["period"], row["output_vat"], row["input_vat"],
            row["vat_payable"], row["closing_input_credit"], row["revenue_projection"],
            row["book_cost_projection"], row["accounting_profit_projection"], row["gate_status"],
        ]
        for row in entities
    ]
    _write_table(ws, headers=overview_headers, rows=overview_rows, money_columns={4,5,6,7,8,9,10}, start_row=13)
    _autosize(ws)

    ws = wb.create_sheet(SHEET_ENTITIES)
    entity_headers = [
        "法人编码", "法人名称", "主体角色类型", "所属期间", "期初留抵", "销项税额", "进项税额",
        "可抵扣进项", "暂不可抵扣进项", "待确认进项", "预缴税款", "实际应纳增值税", "期末留抵",
        "经营投影营业收入", "账面成本投影", "会计利润投影", "门禁状态", "投影状态",
    ]
    entity_values = [
        [
            row["entity_code"], row["entity_name"], row["business_role"], row["period"],
            row["opening_input_credit"], row["output_vat"], row["input_vat"], row["deductible_input_vat"],
            row["nondeductible_input_vat"], row["pending_input_vat"], row["tax_prepayment"],
            row["vat_payable"], row["closing_input_credit"], row["revenue_projection"],
            row["book_cost_projection"], row["accounting_profit_projection"], row["gate_status"],
            row["projection_status"],
        ]
        for row in entities
    ]
    _write_table(ws, headers=entity_headers, rows=entity_values, money_columns=set(range(5,17)))
    _autosize(ws)

    ws = wb.create_sheet(SHEET_PROJECTS)
    project_headers = [
        "项目编号", "项目名称", "归属法人编码", "归属法人名称", "合同总金额", "真实成本", "销项发票金额",
        "进项发票金额", "已支付金额", "四流匹配完整度(%)", "完整度口径", "风险评级", "数据状态",
    ]
    project_values = [
        [
            row["project_code"], row["project_name"], row["entity_code"], row["entity_name"], row["contract_total"],
            row["real_cost"], row["output_invoice_amount"], row["input_invoice_amount"], row["paid_amount"],
            row["four_flow_completeness"], row["four_flow_source"], row["risk_rating"], row["status"],
        ]
        for row in data.get("projects") or []
    ]
    _write_table(ws, headers=project_headers, rows=project_values, money_columns={5,6,7,8,9}, percent_columns={10})
    _autosize(ws)

    ws = wb.create_sheet(SHEET_OUTPUT)
    output_headers = [
        "开票日期", "法人编码", "法人名称", "关联项目编号", "关联项目名称", "关联合同号", "发票号码",
        "买方/交易对手名称", "商品/服务摘要", "不含税金额", "销项税额", "含税总额", "事实状态", "事实ID",
    ]
    output_values = [
        [
            row["invoice_date"], row["entity_code"], row["entity_name"], row["project_code"], row["project_name"],
            row["contract_no"], row["invoice_no"], row["counterparty_name"], row["summary"], row["net_amount"],
            row["vat_amount"], row["gross_amount"], row["fact_status"], row["fact_id"],
        ]
        for row in data.get("output_invoices") or []
    ]
    _write_table(ws, headers=output_headers, rows=output_values, money_columns={10,11,12})
    _autosize(ws)

    ws = wb.create_sheet(SHEET_INPUT)
    input_headers = [
        "开票日期", "法人编码", "法人名称", "关联项目编号", "关联项目名称", "关联合同号", "发票号码",
        "卖方/供应商名称", "商品/服务摘要", "不含税金额", "进项税额", "含税总额", "合规/可抵扣状态",
        "事实状态", "事实ID",
    ]
    input_values = [
        [
            row["invoice_date"], row["entity_code"], row["entity_name"], row["project_code"], row["project_name"],
            row["contract_no"], row["invoice_no"], row["counterparty_name"], row["summary"], row["net_amount"],
            row["vat_amount"], row["gross_amount"], row["deductibility_status"], row["fact_status"], row["fact_id"],
        ]
        for row in data.get("input_invoices") or []
    ]
    _write_table(ws, headers=input_headers, rows=input_values, money_columns={10,11,12})
    _autosize(ws)

    ws = wb.create_sheet(SHEET_PREPAY)
    prepay_headers = [
        "预缴日期", "法人编码", "法人名称", "关联项目编号", "完税凭证号/外部单号", "税务机关", "预缴税种",
        "预缴税额", "备注说明",
    ]
    prepay_values = [
        [
            row["prepayment_date"], row["entity_code"], row["entity_name"], row["project_code"], row["receipt_no"],
            row["tax_authority"], row["tax_type"], row["prepayment_amount"], row["note"],
        ]
        for row in data.get("prepayments") or []
    ]
    _write_table(ws, headers=prepay_headers, rows=prepay_values, money_columns={8})
    _autosize(ws)

    for sheet in wb.worksheets:
        sheet.sheet_view.showGridLines = False
        sheet.row_dimensions[1].height = 24
        for row in sheet.iter_rows():
            for cell in row:
                if cell.row % 2 == 0 and cell.row > 1 and cell.fill.fill_type is None:
                    cell.fill = _SUBTLE_FILL

    stream = BytesIO()
    wb.save(stream)
    return stream.getvalue()


def build_excel_export(
    db,
    *,
    scope: str = "ALL",
    view: str = "current",
    period: str | None = None,
) -> ExcelExportArtifact:
    data = collect_excel_export_data(db, scope=scope, view=view, period=period)
    content = build_excel_workbook(data)
    safe_period = str(data["metadata"]["period"]).replace("/", "-").replace(" ", "_")
    filename = f"税务经营全量导出_{data['metadata']['scope']}_{data['metadata']['view']}_{safe_period}.xlsx"
    return ExcelExportArtifact(content=content, filename=filename)


def content_disposition(filename: str) -> str:
    """RFC 5987-safe attachment header with an ASCII fallback."""
    return f"attachment; filename=tax-export.xlsx; filename*=UTF-8''{quote(filename)}"


__all__ = [
    "ExcelExportArtifact",
    "MONEY_FORMAT",
    "SHEET_NAMES",
    "XLSX_MEDIA_TYPE",
    "build_excel_export",
    "build_excel_workbook",
    "collect_excel_export_data",
    "content_disposition",
]
